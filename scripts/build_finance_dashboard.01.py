from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from datetime import date
from pathlib import Path

OUT=Path("products/GlobalTools-Small-Business-Finance-Dashboard-PRO.xlsx")
OUT.parent.mkdir(parents=True,exist_ok=True)
wb=Workbook(); d=wb.active; d.title="Dashboard"
navy="17324D"; blue="2563EB"; green="15803D"; red="B91C1C"; orange="C2410C"; gray="64748B"; dark="0F172A"; white="FFFFFF"; lg="DCFCE7"; lr="FEE2E2"; lo="FFEDD5"; border=Border(*(Side(style="thin",color="CBD5E1") for _ in range(4)))
cur='$#,##0.00;[Red]-$#,##0.00'; pct='0.0%'; df='mmm d, yyyy'; mf='mmm-yy'
def title(s,t,sub):
 s["A1"]=t;s["A1"].font=Font(size=22,bold=True,color=white);s["A1"].fill=PatternFill("solid",fgColor=navy);s.merge_cells("A1:J1");s.row_dimensions[1].height=34
 s["A2"]=sub;s["A2"].font=Font(size=10,color=gray,italic=True);s.merge_cells("A2:J2")
def head(s,r,n):
 for c in range(1,n+1):
  x=s.cell(r,c);x.font=Font(bold=True,color=white);x.fill=PatternFill("solid",fgColor=blue);x.border=border;x.alignment=Alignment(horizontal="center",wrap_text=True)
def widths(s,n):
 for c in range(1,n+1):
  L=__import__('openpyxl').utils.get_column_letter(c); m=max((len(str(x.value)) for x in s[L] if x.value is not None),default=10);s.column_dimensions[L].width=min(max(m+2,12),32)
def table(s,ref,name):
 t=Table(displayName=name,ref=ref);t.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True,showColumnStripes=False);s.add_table(t)
# Start Here
s=wb.create_sheet("Start Here");title(s,"GlobalTools Finance Dashboard PRO","Professional finance management workbook for small businesses and freelancers.")
rows=[("1","Settings","Set business name, targets, forecast start date and cash reserve."),("2","Transactions","Record income, expenses and cash movements."),("3","Dashboard","Review KPIs and monthly performance."),("4","Cash Flow & Forecast","See daily cash and a 90-day projection."),("5","Goals & KPIs","Track financial targets and progress."),("6","Financial Health","Review liquidity, profitability and risk indicators."),("7","Receivables / Payables","Control money owed to you and money you owe."),("8","Budget / P&L","Plan and compare actual performance.")]
for r,row in enumerate(rows,4):
 for c,v in enumerate(row,1):s.cell(r,c,v).border=border;s.cell(r,c).alignment=Alignment(wrap_text=True,vertical="top")
s["A14"]="Important";s["A14"].font=Font(bold=True,color=red);s["B14"]="This is a management planning tool, not accounting or tax advice.";s.merge_cells("B14:H14");widths(s,8)
# Settings
s=wb.create_sheet("Settings");title(s,"Settings","Customize the workbook before using it.")
vals=[("Business Name","Your Business"),("Owner / Manager","Your Name"),("Base Currency","USD"),("Fiscal Year",2026),("Opening Cash",10000),("Target Monthly Revenue",15000),("Target Net Margin",.20),("Forecast Start Date",date(2026,10,1)),("Cash Reserve Target",5000)]
for r,(a,b) in enumerate(vals,4):s.cell(r,1,a).font=Font(bold=True);s.cell(r,2,b);s.cell(r,1).border=s.cell(r,2).border=border
s["B8"].number_format=s["B9"].number_format=cur;s["B10"].number_format=pct;s["B11"].number_format=df;s["B12"].number_format=cur;widths(s,8)
# Categories
s=wb.create_sheet("Categories");title(s,"Categories","Edit categories used across the workbook.");h=["Category","Type","Group","Active"];s.append([]);s.append(h);head(s,2,4)
for row in [("Sales Revenue","Income","Revenue","Yes"),("Services Revenue","Income","Revenue","Yes"),("Other Income","Income","Other","Yes"),("Payroll","Expense","People","Yes"),("Rent & Utilities","Expense","Operating","Yes"),("Software & Subscriptions","Expense","Operating","Yes"),("Marketing","Expense","Operating","Yes"),("Office & Supplies","Expense","Operating","Yes"),("Travel","Expense","Operating","Yes"),("Professional Services","Expense","Operating","Yes"),("Bank Fees","Expense","Finance","Yes"),("Taxes","Expense","Finance","Yes"),("Equipment","Expense","Capital","Yes"),("Other Expense","Expense","Other","Yes")]:s.append(row)
table(s,f"A2:D{s.max_row}","CategoriesTable");widths(s,4)
# Accounts
s=wb.create_sheet("Accounts");title(s,"Accounts","Bank, cash, credit card and wallet balances.");h=["Account","Type","Opening Balance","Current Balance","Active"];s.append([]);s.append(h);head(s,2,5)
for i,row in enumerate([("Main Checking","Bank",8000,"Yes"),("Business Savings","Bank",2000,"Yes"),("Cash","Cash",500,"Yes"),("Business Credit Card","Credit Card",0,"Yes")],3):
 s.cell(i,1,row[0]);s.cell(i,2,row[1]);s.cell(i,3,row[2]);s.cell(i,5,row[3]);s.cell(i,4,f'=C{i}+SUMIFS(Transactions!$H:$H,Transactions!$G:$G,A{i},Transactions!$E:$E,"Income")-SUMIFS(Transactions!$H:$H,Transactions!$G:$G,A{i},Transactions!$E:$E,"Expense")');s.cell(i,3).number_format=s.cell(i,4).number_format=cur
table(s,f"A2:E{s.max_row}","AccountsTable");widths(s,5)
# Customers
s=wb.create_sheet("Customers");title(s,"Customers","Customer directory.");h=["Customer ID","Customer","Email","Phone","Status","Notes"];s.append([]);s.append(h);head(s,2,6)
for row in [("C-001","Acme Consulting","billing@acme.example","+1 555-0101","Active","Monthly consulting"),("C-002","Northstar Studio","finance@northstar.example","+1 555-0102","Active","Design services"),("C-003","Greenline Retail","accounts@greenline.example","+1 555-0103","Active","Retail account"),("C-004","BrightPath LLC","ap@brightpath.example","+1 555-0104","Active","Project client")]:s.append(row)
table(s,f"A2:F{s.max_row}","CustomersTable");widths(s,6)
# Products
s=wb.create_sheet("Products");title(s,"Products & Services","Maintain your catalog and pricing.");h=["SKU","Product / Service","Type","Unit Price","Cost","Gross Margin","Active"];s.append([]);s.append(h);head(s,2,7)
for i,row in enumerate([("SVC-001","Consulting Package","Service",1500,500),("SVC-002","Monthly Retainer","Service",1200,350),("PRD-001","Digital Template","Product",49,5),("SVC-003","Implementation Project","Service",3000,1200)],3):
 for c,v in enumerate(row,1):s.cell(i,c,v)
 s.cell(i,6,f'=IFERROR((D{i}-E{i})/D{i},0)');s.cell(i,7,"Yes");s.cell(i,4).number_format=s.cell(i,5).number_format=cur;s.cell(i,6).number_format=pct
table(s,f"A2:G{s.max_row}","ProductsTable");widths(s,7)
# Transactions
s=wb.create_sheet("Transactions");title(s,"Transactions","Record all income and expenses using positive amounts.");h=["Date","Reference","Description","Category","Type","Status","Account","Amount","Customer / Vendor","Recurring","Notes"];s.append([]);s.append(h);head(s,2,11)
sample=[(date(2026,1,5),"INV-1001","Consulting Package","Services Revenue","Income","Paid","Main Checking",1500,"Acme Consulting","No","January project"),(date(2026,1,8),"BILL-2001","Software subscription","Software & Subscriptions","Expense","Paid","Main Checking",99,"Cloud Software","Yes","Monthly"),(date(2026,1,12),"INV-1002","Monthly Retainer","Services Revenue","Income","Paid","Main Checking",1200,"Northstar Studio","Yes","Retainer"),(date(2026,1,20),"BILL-2002","Marketing campaign","Marketing","Expense","Paid","Business Credit Card",450,"Ad Platform","No","Campaign"),(date(2026,2,3),"INV-1003","Implementation Project","Services Revenue","Income","Paid","Main Checking",3000,"BrightPath LLC","No","Phase 1"),(date(2026,2,7),"BILL-2003","Professional services","Professional Services","Expense","Paid","Main Checking",600,"Accountant","No","Bookkeeping"),(date(2026,2,15),"INV-1004","Digital Template","Sales Revenue","Income","Paid","Main Checking",49,"Greenline Retail","No","Online sale")]
