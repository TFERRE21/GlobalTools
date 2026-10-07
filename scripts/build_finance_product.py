from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter
from datetime import date

OUT="products/GlobalTools-Small-Business-Finance-Dashboard-PRO.xlsx"
wb=Workbook()

NAVY="17324D"; BLUE="2563EB"; GREEN="15803D"; RED="B91C1C"; ORANGE="C2410C"
LIGHT_GREEN="DCFCE7"; LIGHT_RED="FEE2E2"; LIGHT_ORANGE="FFEDD5"; GRAY="64748B"
LIGHT_GRAY="F1F5F9"; DARK="0F172A"; WHITE="FFFFFF"
CURRENCY='$#,##0.00;[Red]-$#,##0.00'; PCT='0.0%'; DATE='mmm d, yyyy'; MONTH='mmm-yy'
BORDER=Border(left=Side(style="thin",color="CBD5E1"),right=Side(style="thin",color="CBD5E1"),top=Side(style="thin",color="CBD5E1"),bottom=Side(style="thin",color="CBD5E1"))

def title(ws,text,sub):
    ws["A1"]=text; ws["A1"].font=Font(size=22,bold=True,color=WHITE); ws["A1"].fill=PatternFill("solid",fgColor=NAVY)
    ws.merge_cells("A1:J1"); ws.row_dimensions[1].height=34
    ws["A2"]=sub; ws["A2"].font=Font(size=10,color=GRAY,italic=True); ws.merge_cells("A2:J2")
def header(ws,row,n):
    for c in range(1,n+1):
        x=ws.cell(row,c); x.font=Font(bold=True,color=WHITE); x.fill=PatternFill("solid",fgColor=BLUE)
        x.border=BORDER; x.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
def table(ws,ref,name):
    t=Table(displayName=name,ref=ref); t.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True,showColumnStripes=False); ws.add_table(t)
def widths(ws,n):
    for c in range(1,n+1):
        letter=get_column_letter(c); m=max((len(str(v.value)) for v in ws[letter] if v.value is not None),default=10)
        ws.column_dimensions[letter].width=min(max(m+2,12),32)

# Start Here
s=wb.active; s.title="Start Here"; title(s,"GlobalTools Finance Dashboard PRO","Professional small-business finance management workbook.")
items=[
("1","Settings","Set business name, fiscal year, forecast start date and cash reserve."),
("2","Accounts","Add bank, cash and card accounts."),
("3","Transactions","Record every income and expense."),
("4","Budget","Set monthly revenue and expense targets."),
("5","Receivables / Payables","Track invoices, bills, due dates and balances."),
("6","Dashboard","Review KPIs, profitability and cash."),
("7","Cash Flow & Forecast","Review daily cash and the 90-day forecast."),
("8","Goals & KPIs","Measure progress against financial targets."),
("9","Financial Health","Review liquidity, profitability and risk indicators.")]
for r,row in enumerate(items,4):
    for c,v in enumerate(row,1): s.cell(r,c,v); s.cell(r,c).border=BORDER; s.cell(r,c).alignment=Alignment(wrap_text=True,vertical="top")
    s.cell(r,1).font=Font(bold=True,color=BLUE); s.cell(r,2).font=Font(bold=True)
s["A15"]="Important"; s["A15"].font=Font(bold=True,color=RED); s["B15"]="Management tool only; reconcile with bank/accounting records."; s.merge_cells("B15:H15")
widths(s,8)

# Settings
st=wb.create_sheet("Settings"); title(st,"Settings","Customize these inputs before using the workbook.")
settings=[("Business Name","Your Business"),("Owner / Manager","Your Name"),("Base Currency","USD"),("Fiscal Year",2026),("Opening Cash",10000),("Target Monthly Revenue",15000),("Target Net Margin",0.20),("Forecast Start Date",date(2026,10,1)),("Cash Reserve Target",5000)]
for r,(k,v) in enumerate(settings,4):
    st.cell(r,1,k).font=Font(bold=True); st.cell(r,2,v); st.cell(r,1).border=st.cell(r,2).border=BORDER
st["B8"].number_format=st["B9"].number_format=st["B12"].number_format=CURRENCY
st["B10"].number_format=PCT; st["B11"].number_format=DATE
st["A15"]="Tip"; st["A15"].font=Font(bold=True,color=BLUE); st["B15"]="Changing the inputs updates the forecast and KPI calculations."; st.merge_cells("B15:H15")
widths(st,8)

# Categories
cat=wb.create_sheet("Categories"); title(cat,"Categories","Income and expense categories used throughout the workbook.")
cat.append([]); cat.append(["Category","Type","Group","Active"]); header(cat,2,4)
for row in [
("Sales Revenue","Income","Revenue","Yes"),("Services Revenue","Income","Revenue","Yes"),("Other Income","Income","Other","Yes"),
("Payroll","Expense","People","Yes"),("Rent & Utilities","Expense","Operating","Yes"),("Software & Subscriptions","Expense","Operating","Yes"),
("Marketing","Expense","Operating","Yes"),("Office & Supplies","Expense","Operating","Yes"),("Travel","Expense","Operating","Yes"),
("Professional Services","Expense","Operating","Yes"),("Bank Fees","Expense","Finance","Yes"),("Taxes","Expense","Finance","Yes"),
("Equipment","Expense","Capital","Yes"),("Other Expense","Expense","Other","Yes")]: cat.append(row)
table(cat,f"A2:D{cat.max_row}","CategoriesTable"); widths(cat,4)

# Accounts
acc=wb.create_sheet("Accounts"); title(acc,"Accounts","Bank, cash and card balances.")
acc.append([]); acc.append(["Account","Type","Opening Balance","Current Balance","Active"]); header(acc,2,5)
for i,row in enumerate([("Main Checking","Bank",8000,"Yes"),("Business Savings","Bank",2000,"Yes"),("Cash","Cash",500,"Yes"),("Business Credit Card","Credit Card",0,"Yes")],3):
    acc.cell(i,1,row[0]); acc.cell(i,2,row[1]); acc.cell(i,3,row[2]); acc.cell(i,5,row[3])
    acc.cell(i,4,f'=C{i}+SUMIFS(Transactions!$H:$H,Transactions!$G:$G,A{i},Transactions!$E:$E,"Income")-SUMIFS(Transactions!$H:$H,Transactions!$G:$G,A{i},Transactions!$E:$E,"Expense")')
    acc.cell(i,3).number_format=acc.cell(i,4).number_format=CURRENCY
table(acc,f"A2:E{acc.max_row}","AccountsTable"); widths(acc,5)

# Customers
cu=wb.create_sheet("Customers"); title(cu,"Customers","Customer directory.")
cu.append([]); cu.append(["Customer ID","Customer","Email","Phone","Status","Notes"]); header(cu,2,6)
for row in [("C-001","Acme Consulting","billing@acme.example","+1 555-0101","Active","Monthly consulting"),("C-002","Northstar Studio","finance@northstar.example","+1 555-0102","Active","Design services"),("C-003","Greenline Retail","accounts@greenline.example","+1 555-0103","Active","Retail account"),("C-004","BrightPath LLC","ap@brightpath.example","+1 555-0104","Active","Project client")]: cu.append(row)
table(cu,f"A2:F{cu.max_row}","CustomersTable"); widths(cu,6)

# Products
pr=wb.create_sheet("Products"); title(pr,"Products & Services","Maintain products, services and pricing.")
pr.append([]); pr.append(["SKU","Product / Service","Type","Unit Price","Cost","Gross Margin","Active"]); header(pr,2,7)
for i,row in enumerate([("SVC-001","Consulting Package","Service",1500,500),("SVC-002","Monthly Retainer","Service",1200,350),("PRD-001","Digital Template","Product",49,5),("SVC-003","Implementation Project","Service",3000,1200)],3):
    for c,v in enumerate(row,1): pr.cell(i,c,v)
    pr.cell(i,6,f"=IFERROR((D{i}-E{i})/D{i},0)"); pr.cell(i,7,"Yes")
    pr.cell(i,4).number_format=pr.cell(i,5).number_format=CURRENCY; pr.cell(i,6).number_format=PCT
table(pr,f"A2:G{pr.max_row}","ProductsTable"); widths(pr,7)

# Transactions
tx=wb.create_sheet("Transactions"); title(tx,"Transactions","Enter positive amounts and select Income or Expense.")
tx.append([]); tx.append(["Date","Reference","Description","Category","Type","Status","Account","Amount","Customer / Vendor","Recurring","Notes"]); header(tx,2,11)
rows=[
(date(2026,1,5),"INV-1001","Consulting Package","Services Revenue","Income","Paid","Main Checking",1500,"Acme Consulting","No","January project"),
(date(2026,1,8),"BILL-2001","Software subscription","Software & Subscriptions","Expense","Paid","Main Checking",99,"Cloud Software","Yes","Monthly"),
(date(2026,1,12),"INV-1002","Monthly Retainer","Services Revenue","Income","Paid","Main Checking",1200,"Northstar Studio","Yes","Retainer"),
(date(2026,1,20),"BILL-2002","Marketing campaign","Marketing","Expense","Paid","Business Credit Card",450,"Ad Platform","No","Campaign"),
(date(2026,2,3),"INV-1003","Implementation Project","Services Revenue","Income","Paid","Main Checking",3000,"BrightPath LLC","No","Phase 1"),
(date(2026,2,7),"BILL-2003","Professional services","Professional Services","Expense","Paid","Main Checking",600,"Accountant","No","Bookkeeping"),
(date(2026,2,15),"INV-1004","Digital Template","Sales Revenue","Income","Paid","Main Checking",49,"Greenline Retail","No","Online sale")]
for row in rows: tx.append(row)
for r in range(3,tx.max_row+1): tx.cell(r,1).number_format=DATE; tx.cell(r,8).number_format=CURRENCY
table(tx,f"A2:K{tx.max_row}","TransactionsTable"); tx.freeze_panes="A3"; widths(tx,11)

# Budget
bu=wb.create_sheet("Budget"); title(bu,"Budget Planner","Monthly targets vs actual results.")
bu.append([]); bu.append(["Month","Revenue Budget","Expense Budget","Profit Budget","Actual Revenue","Actual Expense","Actual Profit","Revenue Variance","Expense Variance","Profit Variance"]); header(bu,2,10)
for i in range(12):
    r=3+i; bu.cell(r,1,date(2026,i+1,1)); bu.cell(r,1).number_format=MONTH
    bu.cell(r,2,15000); bu.cell(r,3,9000); bu.cell(r,4,f"=B{r}-C{r}")
    bu.cell(r,5,f'=SUMIFS(Transactions!$H:$H,Transactions!$A:$A,">="&A{r},Transactions!$A:$A,"<"&EDATE(A{r},1),Transactions!$E:$E,"Income",Transactions!$F:$F,"<>Cancelled")')
    bu.cell(r,6,f'=SUMIFS(Transactions!$H:$H,Transactions!$A:$A,">="&A{r},Transactions!$A:$A,"<"&EDATE(A{r},1),Transactions!$E:$E,"Expense",Transactions!$F:$F,"<>Cancelled")')
    bu.cell(r,7,f"=E{r}-F{r}"); bu.cell(r,8,f"=E{r}-B{r}"); bu.cell(r,9,f"=F{r}-C{r}"); bu.cell(r,10,f"=G{r}-D{r}")
    for c in range(2,11): bu.cell(r,c).number_format=CURRENCY
table(bu,"A2:J14","BudgetTable"); widths(bu,10)

# Receivables
rc=wb.create_sheet("Receivables"); title(rc,"Accounts Receivable","Track invoices, collections and overdue balances.")
rc.append([]); rc.append(["Invoice","Customer","Issue Date","Due Date","Amount","Paid","Balance","Status","Days Overdue","Notes"]); header(rc,2,10)
for i,row in enumerate([
("INV-1005","Acme Consulting",date(2026,3,1),date(2026,3,31),1800,0,"Open"),
("INV-1006","Northstar Studio",date(2026,3,5),date(2026,3,20),1200,600,"Partial"),
("INV-1007","BrightPath LLC",date(2026,3,10),date(2026,4,10),3000,0,"Open"),
("INV-1008","Greenline Retail",date(2026,3,12),date(2026,3,25),750,750,"Paid")],3):
    inv,cust,issue,due,amt,paid,status=row; rc.cell(i,1,inv); rc.cell(i,2,cust); rc.cell(i,3,issue); rc.cell(i,4,due); rc.cell(i,5,amt); rc.cell(i,6,paid); rc.cell(i,7,f"=E{i}-F{i}"); rc.cell(i,8,status); rc.cell(i,9,f'=IF(OR(G{i}=0,H{i}="Paid"),0,MAX(0,TODAY()-D{i}))')
    rc.cell(i,3).number_format=rc.cell(i,4).number_format=DATE
    for c in [5,6,7]: rc.cell(i,c).number_format=CURRENCY
table(rc,f"A2:J{rc.max_row}","ReceivablesTable"); widths(rc,10)

# Payables
pa=wb.create_sheet("Payables"); title(pa,"Accounts Payable","Track bills, due dates and obligations.")
pa.append([]); pa.append(["Bill","Vendor","Issue Date","Due Date","Amount","Paid","Balance","Status","Days Overdue","Priority"]); header(pa,2,10)
for i,row in enumerate([
("BILL-3001","Cloud Software",date(2026,3,1),date(2026,3,15),99,99,"Paid","Normal"),
("BILL-3002","Ad Platform",date(2026,3,5),date(2026,3,18),450,0,"Open","High"),
("BILL-3003","Accountant",date(2026,3,10),date(2026,3,30),600,0,"Open","Medium"),
("BILL-3004","Office Supplier",date(2026,3,12),date(2026,4,5),320,0,"Open","Normal")],3):
    bill,vendor,issue,due,amt,paid,status,priority=row; pa.cell(i,1,bill); pa.cell(i,2,vendor); pa.cell(i,3,issue); pa.cell(i,4,due); pa.cell(i,5,amt); pa.cell(i,6,paid); pa.cell(i,7,f"=E{i}-F{i}"); pa.cell(i,8,status); pa.cell(i,9,f'=IF(OR(G{i}=0,H{i}="Paid"),0,MAX(0,TODAY()-D{i}))'); pa.cell(i,10,priority)
    pa.cell(i,3).number_format=pa.cell(i,4).number_format=DATE
    for c in [5,6,7]: pa.cell(i,c).number_format=CURRENCY
table(pa,f"A2:J{pa.max_row}","PayablesTable"); widths(pa,10)

# Monthly Summary
ms=wb.create_sheet("Monthly Summary"); title(ms,"Monthly Financial Summary","Revenue, expenses, profit, margin and budget variance.")
ms.append([]); ms.append(["Month","Revenue","Expenses","Net Profit","Net Margin","Revenue Budget","Expense Budget","Profit Budget","Revenue Var.","Expense Var.","Profit Var.","Cumulative Profit"]); header(ms,2,12)
for i in range(12):
    r=3+i; ms.cell(r,1,date(2026,i+1,1)); ms.cell(r,1).number_format=MONTH
    ms.cell(r,2,f'=SUMIFS(Transactions!$H:$H,Transactions!$A:$A,">="&A{r},Transactions!$A:$A,"<"&EDATE(A{r},1),Transactions!$E:$E,"Income",Transactions!$F:$F,"<>Cancelled")')
    ms.cell(r,3,f'=SUMIFS(Transactions!$H:$H,Transactions!$A:$A,">="&A{r},Transactions!$A:$A,"<"&EDATE(A{r},1),Transactions!$E:$E,"Expense",Transactions!$F:$F,"<>Cancelled")')
    ms.cell(r,4,f"=B{r}-C{r}"); ms.cell(r,5,f"=IFERROR(D{r}/B{r},0)"); ms.cell(r,6,f"=Budget!B{r}"); ms.cell(r,7,f"=Budget!C{r}"); ms.cell(r,8,f"=Budget!D{r}"); ms.cell(r,9,f"=B{r}-F{r}"); ms.cell(r,10,f"=C{r}-G{r}"); ms.cell(r,11,f"=D{r}-H{r}"); ms.cell(r,12,f"=SUM($D$3:D{r})")
    for c in [2,3,4,6,7,8,9,10,11,12]: ms.cell(r,c).number_format=CURRENCY
    ms.cell(r,5).number_format=PCT
table(ms,"A2:L14","MonthlySummaryTable"); widths(ms,12)

# P&L
pl=wb.create_sheet("P&L Statement"); title(pl,"Profit & Loss Statement","Monthly management P&L.")
pl.append([]); pl.append(["Line Item","Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec","YTD"]); header(pl,2,14)
for r,(label,typ) in enumerate([("Revenue","Income"),("Operating Expenses","Expense"),("Net Profit","Profit")],3):
    pl.cell(r,1,label)
    for m in range(1,13):
        c=m+1; col=get_column_letter(c)
        if typ=="Profit": f=f"={col}3-{col}4"
        else: f=f'=SUMIFS(Transactions!$H:$H,Transactions!$A:$A,">="&DATE(2026,{m},1),Transactions!$A:$A,"<"&EDATE(DATE(2026,{m},1),1),Transactions!$E:$E,"{typ}",Transactions!$F:$F,"<>Cancelled")'
        pl.cell(r,c,f); pl.cell(r,c).number_format=CURRENCY
    pl.cell(r,14,f"=SUM(B{r}:M{r})"); pl.cell(r,14).number_format=CURRENCY
pl["A6"]="Net Margin"; pl["A6"].font=Font(bold=True,color=NAVY)
for c in range(2,15):
    col=get_column_letter(c); pl.cell(6,c,f"=IFERROR({col}5/{col}3,0)"); pl.cell(6,c).number_format=PCT
widths(pl,14)

# Dashboard
d=wb.create_sheet("Dashboard"); title(d,"GlobalTools Finance Dashboard PRO","Executive view with profitability, cash, targets and risk.")
cards=[("B4","Revenue YTD",'=SUMIFS(Transactions!$H:$H,Transactions!$E:$E,"Income",Transactions!$F:$F,"<>Cancelled")',GREEN,CURRENCY),
("D4","Expenses YTD",'=SUMIFS(Transactions!$H:$H,Transactions!$E:$E,"Expense",Transactions!$F:$F,"<>Cancelled")',RED,CURRENCY),
("F4","Net Profit","=B5-D5",BLUE,CURRENCY),("H4","Net Margin","=IFERROR(F5/B5,0)",ORANGE,PCT)]
for cell,label,formula,color,fmt in cards:
    c=d[cell].column;r=d[cell].row; d.cell(r,c,label); d.cell(r,c).font=Font(bold=True,color=WHITE); d.cell(r,c).fill=PatternFill("solid",fgColor=color); d.merge_cells(start_row=r,start_column=c,end_row=r,end_column=c+1)
    d.cell(r+1,c,formula); d.cell(r+1,c).font=Font(size=16,bold=True,color=DARK); d.cell(r+1,c).fill=PatternFill("solid",fgColor=LIGHT_GRAY); d.cell(r+1,c).number_format=fmt; d.merge_cells(start_row=r+1,start_column=c,end_row=r+1,end_column=c+1)
for cell,label,formula,fmt in [("B8","Cash Balance","=SUM(Accounts!D3:D100)",CURRENCY),("D8","Receivables","=SUM(Receivables!G3:G100)",CURRENCY),("F8","Payables","=SUM(Payables!G3:G100)",CURRENCY),("H8","Financial Health","='Financial Health'!B13","0")]:
    c=d[cell].column;r=d[cell].row; d.cell(r,c,label).font=Font(bold=True,color=GRAY); d.merge_cells(start_row=r,start_column=c,end_row=r,end_column=c+1); d.cell(r+1,c,formula).font=Font(size=14,bold=True,color=DARK); d.cell(r+1,c).number_format=fmt; d.merge_cells(start_row=r+1,start_column=c,end_row=r+1,end_column=c+1)
d["B12"]="Monthly Performance"; d["B12"].font=Font(size=14,bold=True,color=NAVY)
for c,h in enumerate(["Month","Revenue","Expenses","Profit","Margin"],2): d.cell(13,c,h)
header(d,13,6)
for i in range(12):
    r=14+i; d.cell(r,2,f"='Monthly Summary'!A{3+i}"); d.cell(r,2).number_format=MONTH; d.cell(r,3,f"='Monthly Summary'!B{3+i}"); d.cell(r,4,f"='Monthly Summary'!C{3+i}"); d.cell(r,5,f"='Monthly Summary'!D{3+i}"); d.cell(r,6,f"='Monthly Summary'!E{3+i}")
    for c in [3,4,5]: d.cell(r,c).number_format=CURRENCY
    d.cell(r,6).number_format=PCT
bar=BarChart(); bar.title="Revenue vs Expenses"; bar.height=7; bar.width=13; bar.add_data(Reference(d,min_col=3,max_col=4,min_row=13,max_row=25),titles_from_data=True); bar.set_categories(Reference(d,min_col=2,min_row=14,max_row=25)); d.add_chart(bar,"H12")
d["B28"]="PRO modules"; d["B28"].font=Font(size=14,bold=True,color=NAVY)
for cell,lab,desc in [("B29","Daily Cash Flow","Actual vs projected cash movement."),("D29","90-Day Forecast","See low-cash periods before they happen."),("F29","Goals & KPIs","Measure financial targets and progress."),("H29","Financial Health","Spot liquidity, margin and budget risks.")]:
    c=d[cell].column;r=d[cell].row; d.cell(r,c,lab).font=Font(bold=True,color=BLUE); d.cell(r+1,c,desc).font=Font(size=9,color=GRAY); d.merge_cells(start_row=r,start_column=c,end_row=r,end_column=c+1); d.merge_cells(start_row=r+1,start_column=c,end_row=r+1,end_column=c+1)
widths(d,10)

# Cash Flow & Forecast
cf=wb.create_sheet("Cash Flow & Forecast"); title(cf,"Cash Flow & Forecast","Daily cash visibility plus a 90-day forward projection.")
cf.append([]); cf.append(["Date","Opening Cash","Actual Inflows","Actual Outflows","Net Actual","Scheduled AR","Scheduled AP","Projected Net","Projected Closing Cash","Reserve Gap"]); header(cf,2,10)
for i in range(91):
    r=3+i
    cf.cell(r,1,"=Settings!$B$11" if i==0 else f"=A{r-1}+1"); cf.cell(r,1).number_format=DATE
    cf.cell(r,2,"=SUM(Accounts!$D$3:$D$100)" if i==0 else f"=I{r-1}")
    cf.cell(r,3,f'=SUMIFS(Transactions!$H:$H,Transactions!$A:$A,A{r},Transactions!$E:$E,"Income",Transactions!$F:$F,"<>Cancelled")')
    cf.cell(r,4,f'=SUMIFS(Transactions!$H:$H,Transactions!$A:$A,A{r},Transactions!$E:$E,"Expense",Transactions!$F:$F,"<>Cancelled")')
    cf.cell(r,5,f"=C{r}-D{r}"); cf.cell(r,6,f'=SUMIFS(Receivables!$G:$G,Receivables!$D:$D,A{r},Receivables!$H:$H,"<>Paid")'); cf.cell(r,7,f'=SUMIFS(Payables!$G:$G,Payables!$D:$D,A{r},Payables!$H:$H,"<>Paid")')
    cf.cell(r,8,f'=F{r}-G{r}+(AVERAGE(\'Monthly Summary\'!$D$3:$D$14)/30)'); cf.cell(r,9,f"=B{r}+H{r}"); cf.cell(r,10,f"=MAX(0,Settings!$B$12-I{r})")
    for c in range(2,11): cf.cell(r,c).number_format=CURRENCY
table(cf,"A2:J93","CashForecastTable"); cf.freeze_panes="A3"
cf["L4"]="Forecast KPIs"; cf["L4"].font=Font(size=14,bold=True,color=NAVY)
for r,(lab,form,fmt) in enumerate([("Starting Cash","=B3",CURRENCY),("90-Day Ending Cash","=I93",CURRENCY),("Lowest Cash","=MIN(I3:I93)",CURRENCY),("Lowest Cash Date","=INDEX(A3:A93,MATCH(M7,I3:I93,0))",DATE),("Reserve Target","=Settings!B12",CURRENCY),("Days Below Reserve",'=COUNTIF(I3:I93,"<"&Settings!B12)',"0")],5):
    cf.cell(r,12,lab).font=Font(bold=True,color=GRAY); cf.cell(r,13,form); cf.cell(r,13).number_format=fmt
line=LineChart(); line.title="90-Day Projected Closing Cash"; line.height=8; line.width=15; line.add_data(Reference(cf,min_col=9,min_row=2,max_row=93),titles_from_data=True); line.set_categories(Reference(cf,min_col=1,min_row=3,max_row=93)); cf.add_chart(line,"L13")
widths(cf,13)

# Goals & KPIs
go=wb.create_sheet("Goals & KPIs"); title(go,"Goals & KPIs","Track financial targets and monthly progress.")
go.append([]); go.append(["Goal","Target","Actual","Progress","Status","Action / Owner"]); header(go,2,6)
goals=[("Monthly Revenue","=Settings!B9",'=SUMIFS(Transactions!$H:$H,Transactions!$E:$E,"Income",Transactions!$F:$F,"<>Cancelled")'),
("Net Margin","=Settings!B10",'=IFERROR((SUMIFS(Transactions!$H:$H,Transactions!$E:$E,"Income")-SUMIFS(Transactions!$H:$H,Transactions!$E:$E,"Expense"))/SUMIFS(Transactions!$H:$H,Transactions!$E:$E,"Income"),0)'),
("Cash Reserve","=Settings!B12","=SUM(Accounts!D3:D100)"),("Receivables Collected","=SUM(Receivables!E3:E100)","=SUM(Receivables!F3:F100)"),("Budget Profit","=SUM(Budget!D3:D14)","=SUM('Monthly Summary'!D3:D14)")]
for r,(g,t,a) in enumerate(goals,3):
    go.cell(r,1,g); go.cell(r,2,t); go.cell(r,3,a); go.cell(r,4,f"=MIN(1,IFERROR(C{r}/B{r},0))"); go.cell(r,5,f'=IF(D{r}>=1,"On Track",IF(D{r}>=0.75,"Watch","Needs Attention"))'); go.cell(r,6,"Review weekly")
    if g=="Net Margin": go.cell(r,2).number_format=go.cell(r,3).number_format=PCT
    else: go.cell(r,2).number_format=go.cell(r,3).number_format=CURRENCY
    go.cell(r,4).number_format=PCT
table(go,"A2:F7","GoalsTable"); widths(go,6)

# Financial Health
fh=wb.create_sheet("Financial Health"); title(fh,"Financial Health Score","Management scorecard for profitability, liquidity, budget and forecast risk.")
fh.append([]); fh.append(["Indicator","Healthy Range","Actual","Score","Status","Why it matters"]); header(fh,2,6)
inds=[
("Net Margin",">= 20%","=IFERROR('Monthly Summary'!D14/'Monthly Summary'!B14,0)","=MIN(100,MAX(0,C3/Settings!B10*100))","Profitability"),
("Expense Ratio","<= 80%","=IFERROR('Monthly Summary'!C14/'Monthly Summary'!B14,0)","=MIN(100,MAX(0,(1-C4)/0.2*100))","Cost control"),
("Cash Reserve","Cash >= target","=SUM(Accounts!D3:D100)","=MIN(100,IFERROR(C5/Settings!B12*100,0))","Liquidity"),
("Receivables","Lower is better","=SUM(Receivables!G3:G100)","=MAX(0,100-MIN(100,C6/Settings!B12*100))","Collections"),
("Payables","Manage obligations","=SUM(Payables!G3:G100)","=MAX(0,100-MIN(100,C7/Settings!B12*100))","Obligations"),
("Revenue vs Budget",">= 100%","=IFERROR(SUM('Monthly Summary'!B3:B14)/SUM('Monthly Summary'!F3:F14),0)","=MIN(100,C8*100)","Growth"),
("Profit vs Budget",">= 100%","=IFERROR(SUM('Monthly Summary'!D3:D14)/SUM('Monthly Summary'!H3:H14),0)","=MIN(100,C9*100)","Execution"),
("Cash Forecast","No days below reserve","=COUNTIF('Cash Flow & Forecast'!I3:I93,"<"&Settings!B12)","=MAX(0,100-C10*5)","Forward risk")]
for r,(name,target,actual,score,why) in enumerate(inds,3):
    fh.cell(r,1,name); fh.cell(r,2,target); fh.cell(r,3,actual); fh.cell(r,4,score); fh.cell(r,5,f'=IF(D{r}>=80,"Healthy",IF(D{r}>=60,"Watch","Risk"))'); fh.cell(r,6,why)
    if name in ["Net Margin","Expense Ratio","Revenue vs Budget","Profit vs Budget"]: fh.cell(r,3).number_format=PCT
    elif name!="Cash Forecast": fh.cell(r,3).number_format=CURRENCY
    fh.cell(r,4).number_format="0"
table(fh,"A2:F10","HealthIndicatorsTable")
fh["A13"]="Overall Financial Health Score"; fh["A13"].font=Font(size=16,bold=True,color=NAVY); fh["B13"]="=ROUND(AVERAGE(D3:D10),0)"; fh["B13"].font=Font(size=24,bold=True,color=BLUE); fh["C13"]='=IF(B13>=80,"HEALTHY",IF(B13>=60,"WATCH","ACTION REQUIRED"))'; fh["C13"].font=Font(size=14,bold=True)
fh["A15"]="Important"; fh["B15"]="Management indicator only — not a credit rating, accounting opinion or financial advice."; fh.merge_cells("B15:F15")
widths(fh,6)

# Final formatting
for ws in wb.worksheets:
    ws.sheet_view.showGridLines=False
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None: cell.alignment=Alignment(vertical="center",wrap_text=True)
    ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
wb.properties.title="GlobalTools Small Business Finance Dashboard PRO"
wb.properties.creator="GlobalTools"
wb.properties.description="Professional finance workbook with dashboard, budget, P&L, receivables, payables, daily cash flow, 90-day forecast, goals and financial health."
wb.save(OUT)
print(OUT)
